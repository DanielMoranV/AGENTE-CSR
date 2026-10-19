"""
Generador de Reportes Ejecutivos en Word (.docx) para AGENTE-CSR.
Diseñado con tipografía institucional, cajas de resumen ejecutivo,
tablas formateadas con estilo corporativo y auditoría de seguridad.
"""

from datetime import datetime
from pathlib import Path
from typing import Optional, Dict, List, Any
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from src.reportes.estilos import (
    RGB_NAVY_PRIMARY,
    RGB_BLUE_SECONDARY,
    RGB_GRAY_BG,
    RGB_DARK_TEXT,
    HEX_NAVY_PRIMARY,
    HEX_BLUE_SECONDARY,
    HEX_GRAY_BG,
    HEX_BORDER,
    FONT_PRIMARY_DOCX,
    NOMBRE_INSTITUCION,
    CONFIDENCIALIDAD_TEXTO
)
from src.reportes.validaciones import validar_contenido_documento, validar_destino_reporte
from src.security.audit import registrar_evento_auditoria


def _aplicar_sombreado_celda(cell, hex_color: str):
    """Aplica color de fondo a una celda de tabla en python-docx."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))


def _aplicar_borde_celda(cell, hex_color: str = HEX_BORDER):
    """Aplica bordes finos a una celda."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
        <w:left w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
        <w:right w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
    </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))


def generar_reporte_word(
    titulo: str,
    resumen: str,
    tablas: Optional[List[Dict[str, Any]]],
    conclusiones: Optional[List[str]],
    nombre_archivo: str,
    subtitulo: str = "",
    usuario: Optional[str] = None
) -> Path:
    """
    Genera un informe institucional en formato Microsoft Word (.docx).
    
    Args:
        titulo: Título principal del reporte ejecutivo.
        resumen: Texto detallado del resumen ejecutivo.
        tablas: Lista de diccionarios con {'nombre': str, 'datos': pd.DataFrame, 'formatos': dict}.
        conclusiones: Lista de viñetas con los hallazgos y recomendaciones clave.
        nombre_archivo: Nombre del archivo de salida (ej: 'informe_facturacion.docx').
        subtitulo: Subtítulo con rango de fechas o contexto de análisis.
        usuario: Usuario de la sesión para auditoría.
        
    Returns:
        Path resuelto del archivo .docx generado en workspace/resultados/word/.
    """
    # 1. Validaciones previas de seguridad y contenido
    validar_contenido_documento(titulo, resumen, tablas, conclusiones)
    ruta_destino = validar_destino_reporte(nombre_archivo, "docx")

    doc = Document()

    # Configuración de márgenes estándar (2.0 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Encabezado institucional
        header_p = section.header.paragraphs[0]
        header_p.text = f"{NOMBRE_INSTITUCION} — Dirección y Gerencia General"
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_p.runs[0].font.size = Pt(8.5)
        header_p.runs[0].font.color.rgb = RGBColor(120, 120, 120)
        
        # Pie de página con confidencialidad
        footer_p = section.footer.paragraphs[0]
        footer_p.text = f"{CONFIDENCIALIDAD_TEXTO} | Emisión: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_p.runs[0].font.size = Pt(8)
        footer_p.runs[0].font.color.rgb = RGBColor(140, 140, 140)

    # 2. Título y Subtítulo
    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_titulo = p_titulo.add_run(titulo)
    run_titulo.font.name = FONT_PRIMARY_DOCX
    run_titulo.font.size = Pt(18)
    run_titulo.font.bold = True
    run_titulo.font.color.rgb = RGBColor(*RGB_NAVY_PRIMARY)

    if subtitulo:
        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run(subtitulo)
        run_sub.font.name = FONT_PRIMARY_DOCX
        run_sub.font.size = Pt(10.5)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3. Sección: Resumen Ejecutivo
    h1_resumen = doc.add_paragraph()
    r_h1 = h1_resumen.add_run("1. Resumen Ejecutivo")
    r_h1.font.name = FONT_PRIMARY_DOCX
    r_h1.font.size = Pt(13)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(*RGB_BLUE_SECONDARY)
    h1_resumen.paragraph_format.space_after = Pt(4)

    # Caja sombreada para el resumen
    tabla_resumen = doc.add_table(rows=1, cols=1)
    tabla_resumen.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda_resumen = tabla_resumen.rows[0].cells[0]
    _aplicar_sombreado_celda(celda_resumen, HEX_GRAY_BG)
    celda_resumen.paragraphs[0].text = resumen
    celda_resumen.paragraphs[0].runs[0].font.name = FONT_PRIMARY_DOCX
    celda_resumen.paragraphs[0].runs[0].font.size = Pt(10)
    celda_resumen.paragraphs[0].runs[0].font.color.rgb = RGBColor(*RGB_DARK_TEXT)
    _aplicar_borde_celda(celda_resumen, HEX_BLUE_SECONDARY)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 4. Sección: Tablas de Datos
    if tablas:
        h1_tablas = doc.add_paragraph()
        r_h2 = h1_tablas.add_run("2. Información Detallada y Cuadros Analíticos")
        r_h2.font.name = FONT_PRIMARY_DOCX
        r_h2.font.size = Pt(13)
        r_h2.font.bold = True
        r_h2.font.color.rgb = RGBColor(*RGB_BLUE_SECONDARY)
        h1_tablas.paragraph_format.space_after = Pt(6)

        for item in tablas:
            nombre_tabla = item.get("nombre", "Cuadro Analítico")
            df = item.get("datos")
            if df is None or not isinstance(df, pd.DataFrame) or df.empty:
                continue

            p_nom = doc.add_paragraph()
            r_nom = p_nom.add_run(f"Cuadro: {nombre_tabla}")
            r_nom.font.bold = True
            r_nom.font.size = Pt(10.5)
            r_nom.font.color.rgb = RGBColor(*RGB_NAVY_PRIMARY)

            # Construir tabla
            cols = list(df.columns)
            tabla_docx = doc.add_table(rows=len(df) + 1, cols=len(cols))
            tabla_docx.alignment = WD_TABLE_ALIGNMENT.CENTER

            # Encabezado
            hdr_cells = tabla_docx.rows[0].cells
            for idx, col_name in enumerate(cols):
                hdr_cells[idx].text = str(col_name).replace("_", " ").title()
                _aplicar_sombreado_celda(hdr_cells[idx], HEX_NAVY_PRIMARY)
                _aplicar_borde_celda(hdr_cells[idx], HEX_NAVY_PRIMARY)
                p = hdr_cells[idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.runs[0].font.name = FONT_PRIMARY_DOCX
                p.runs[0].font.size = Pt(9.5)
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

            # Filas de datos
            formatos = item.get("formatos", {})
            for r_idx, (_, row) in enumerate(df.iterrows(), start=1):
                row_cells = tabla_docx.rows[r_idx].cells
                fill_color = HEX_GRAY_BG if (r_idx % 2 == 0) else "FFFFFF"

                for c_idx, col_name in enumerate(cols):
                    val = row[col_name]
                    # Formatear valor
                    fmt = formatos.get(col_name, "")
                    if fmt == "moneda" and isinstance(val, (int, float)):
                        texto_val = f"S/ {val:,.2f}"
                        align = WD_ALIGN_PARAGRAPH.RIGHT
                    elif fmt == "porcentaje" and isinstance(val, (int, float)):
                        texto_val = f"{val * 100:.1f}%"
                        align = WD_ALIGN_PARAGRAPH.RIGHT
                    elif isinstance(val, (int, float)):
                        texto_val = f"{val:,.2f}" if isinstance(val, float) else f"{val:,}"
                        align = WD_ALIGN_PARAGRAPH.RIGHT
                    else:
                        texto_val = str(val or "")
                        align = WD_ALIGN_PARAGRAPH.LEFT

                    row_cells[c_idx].text = texto_val
                    _aplicar_sombreado_celda(row_cells[c_idx], fill_color)
                    _aplicar_borde_celda(row_cells[c_idx], HEX_BORDER)
                    p = row_cells[c_idx].paragraphs[0]
                    p.alignment = align
                    if p.runs:
                        p.runs[0].font.name = FONT_PRIMARY_DOCX
                        p.runs[0].font.size = Pt(9)
                        p.runs[0].font.color.rgb = RGBColor(*RGB_DARK_TEXT)

            doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 5. Sección: Conclusiones y Recomendaciones
    if conclusiones:
        h1_conc = doc.add_paragraph()
        r_h3 = h1_conc.add_run("3. Conclusiones y Recomendaciones Estratégicas")
        r_h3.font.name = FONT_PRIMARY_DOCX
        r_h3.font.size = Pt(13)
        r_h3.font.bold = True
        r_h3.font.color.rgb = RGBColor(*RGB_BLUE_SECONDARY)
        h1_conc.paragraph_format.space_after = Pt(4)

        for conc in conclusiones:
            p_bullet = doc.add_paragraph(style='List Bullet')
            run_bullet = p_bullet.add_run(conc)
            run_bullet.font.name = FONT_PRIMARY_DOCX
            run_bullet.font.size = Pt(10)
            run_bullet.font.color.rgb = RGBColor(*RGB_DARK_TEXT)

    # Guardar documento
    doc.save(str(ruta_destino))

    # Auditoría
    registrar_evento_auditoria(
        accion="GENERACION_WORD",
        modulo="reportes.word",
        operacion="generar_reporte_word",
        detalles={
            "archivo": str(ruta_destino),
            "titulo": titulo[:80],
            "tablas_incluidas": len(tablas) if tablas else 0,
            "conclusiones_count": len(conclusiones) if conclusiones else 0
        },
        usuario=usuario,
        estado="EXITO"
    )

    return ruta_destino
