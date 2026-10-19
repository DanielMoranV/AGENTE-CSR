"""
Generador de Reportes Ejecutivos en Excel (.xlsx) para AGENTE-CSR.
Diseñado con formato corporativo, estilos de alta legibilidad para Gerencia,
protección contra inyección de fórmulas y validación estricta de rutas.
"""

from datetime import datetime
from pathlib import Path
from typing import Optional, Dict, List, Any
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

from src.security.permissions import validar_ruta_escritura, sanitizar_valor_excel, RESULTADOS_DIR
from src.security.audit import registrar_evento_auditoria


# Paleta corporativa institucional
COLOR_HEADER_BG = "1F4E78"       # Azul institucional oscuro
COLOR_HEADER_FG = "FFFFFF"       # Blanco
COLOR_ZEBRA_BG = "F9FAFC"        # Fondo suave alternado
COLOR_TOTAL_BG = "D9E1F2"        # Azul claro para totales
COLOR_BORDER = "D3D3D3"          # Gris suave para líneas de celda


def generar_reporte_excel(
    df: pd.DataFrame,
    nombre_archivo: str,
    titulo: str,
    subtitulo: str = "",
    formatos_columnas: Optional[Dict[str, str]] = None,
    incluir_totales: bool = True,
    columnas_totales: Optional[List[str]] = None,
    crear_grafico: bool = False,
    grafico_config: Optional[Dict[str, Any]] = None,
    usuario: str = "Gerencia"
) -> Path:
    """
    Genera un reporte Excel con estilo ejecutivo profesional y lo guarda en workspace/resultados/excel/.
    
    Args:
        df: DataFrame con los datos a exportar.
        nombre_archivo: Nombre del archivo .xlsx (ej: 'facturacion_agosto_septiembre.xlsx').
        titulo: Título visible en el encabezado del reporte.
        subtitulo: Subtítulo con rango de fechas, filtros o contexto.
        formatos_columnas: Dict con {columna: formato} ('moneda', 'porcentaje', 'entero', 'decimal').
        incluir_totales: Si se debe agregar fila de Total al final.
        columnas_totales: Lista de nombres de columnas a sumar (si es None, suma numéricas).
        crear_grafico: Si se debe incrustar un gráfico comparativo.
        grafico_config: Configuración del gráfico {'categoria': 'col_x', 'valores': ['col_y1', 'col_y2'], 'titulo': '...'}.
        usuario: Usuario solicitante para auditoría.
        
    Returns:
        Path resuelto y validado del archivo creado.
    """
    if not nombre_archivo.endswith(".xlsx"):
        nombre_archivo += ".xlsx"

    # Ruta validada estrictamente en workspace/resultados/excel/
    ruta_destino = validar_ruta_escritura(
        RESULTADOS_DIR / "excel" / nombre_archivo
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte Ejecutivo"
    ws.views.sheetView[0].showGridLines = True

    # Estilos base
    font_titulo = Font(name="Segoe UI", size=15, bold=True, color="1F4E78")
    font_subtitulo = Font(name="Segoe UI", size=10, italic=True, color="595959")
    font_header = Font(name="Segoe UI", size=11, bold=True, color=COLOR_HEADER_FG)
    font_datos = Font(name="Segoe UI", size=10)
    font_total = Font(name="Segoe UI", size=11, bold=True, color="000000")

    fill_header = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
    fill_zebra = PatternFill(start_color=COLOR_ZEBRA_BG, end_color=COLOR_ZEBRA_BG, fill_type="solid")
    fill_total = PatternFill(start_color=COLOR_TOTAL_BG, end_color=COLOR_TOTAL_BG, fill_type="solid")

    border_thin = Border(
        left=Side(style="thin", color=COLOR_BORDER),
        right=Side(style="thin", color=COLOR_BORDER),
        top=Side(style="thin", color=COLOR_BORDER),
        bottom=Side(style="thin", color=COLOR_BORDER)
    )
    border_total = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000")
    )

    # 1. Cabecera Institucional
    ws.cell(row=1, column=1, value=titulo).font = font_titulo
    sub_texto = subtitulo or f"Generado el {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws.cell(row=2, column=1, value=sub_texto).font = font_subtitulo

    fila_inicio_tabla = 4
    columnas = list(df.columns)
    formatos_columnas = formatos_columnas or {}

    # 2. Encabezados de Columna
    for col_idx, col_name in enumerate(columnas, start=1):
        cell = ws.cell(row=fila_inicio_tabla, column=col_idx, value=str(col_name).replace("_", " ").title())
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    ws.row_dimensions[fila_inicio_tabla].height = 28

    # 3. Filas de Datos
    fila_actual = fila_inicio_tabla + 1
    for row_idx, row in df.iterrows():
        fill_actual = fill_zebra if (row_idx % 2 == 1) else None
        
        for col_idx, col_name in enumerate(columnas, start=1):
            val = row[col_name]
            val_seguro = sanitizar_valor_excel(val)
            cell = ws.cell(row=fila_actual, column=col_idx, value=val_seguro)
            cell.font = font_datos
            cell.border = border_thin
            if fill_actual:
                cell.fill = fill_actual

            # Aplicar formato de número según tipo configurado o inferido
            tipo_fmt = formatos_columnas.get(col_name, "")
            if tipo_fmt == "moneda":
                cell.number_format = '"S/" #,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif tipo_fmt == "porcentaje":
                cell.number_format = '0.0%'
                cell.alignment = Alignment(horizontal="right")
            elif tipo_fmt == "entero":
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right")
            elif tipo_fmt == "decimal":
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")

        fila_actual += 1

    # 4. Fila de Totales (opcional)
    if incluir_totales and len(df) > 0:
        ws.cell(row=fila_actual, column=1, value="TOTAL GENERAL").font = font_total
        ws.cell(row=fila_actual, column=1).fill = fill_total
        ws.cell(row=fila_actual, column=1).border = border_total

        # Determinar qué columnas sumar
        if columnas_totales is None:
            cols_a_sumar = [c for c in columnas if pd.api.types.is_numeric_dtype(df[c]) and formatos_columnas.get(c) != "porcentaje"]
        else:
            cols_a_sumar = columnas_totales

        for col_idx, col_name in enumerate(columnas, start=1):
            cell = ws.cell(row=fila_actual, column=col_idx)
            cell.border = border_total
            cell.fill = fill_total

            if col_name in cols_a_sumar:
                col_letter = get_column_letter(col_idx)
                cell.value = f"=SUM({col_letter}{fila_inicio_tabla+1}:{col_letter}{fila_actual-1})"
                cell.font = font_total
                
                tipo_fmt = formatos_columnas.get(col_name, "")
                if tipo_fmt == "moneda":
                    cell.number_format = '"S/" #,##0.00'
                elif tipo_fmt == "entero":
                    cell.number_format = '#,##0'
                else:
                    cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")

        fila_actual += 1

    # 5. Autoajustar ancho de columnas
    for col_idx, col_name in enumerate(columnas, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(col_name)),
            max([len(str(df.iloc[r][col_name])) for r in range(len(df))] or [0])
        )
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # 6. Gráfico opcional
    if crear_grafico and grafico_config and len(df) > 0:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = grafico_config.get("titulo", "Comparativo")
        chart.height = 14
        chart.width = 22

        cat_col = grafico_config.get("categoria", columnas[0])
        cat_idx = columnas.index(cat_col) + 1
        val_cols = grafico_config.get("valores", [columnas[1]])

        for v_col in val_cols:
            if v_col in columnas:
                v_idx = columnas.index(v_col) + 1
                data_ref = Reference(ws, min_col=v_idx, min_row=fila_inicio_tabla, max_row=fila_actual - (2 if incluir_totales else 1))
                chart.add_data(data_ref, titles_from_data=True)

        cats_ref = Reference(ws, min_col=cat_idx, min_row=fila_inicio_tabla + 1, max_row=fila_actual - (2 if incluir_totales else 1))
        chart.set_categories(cats_ref)

        chart_pos = f"{get_column_letter(len(columnas) + 2)}4"
        ws.add_chart(chart, chart_pos)

    # Guardar
    wb.save(ruta_destino)

    registrar_evento_auditoria(
        accion="GENERACION_EXCEL",
        modulo="reportes.excel",
        operacion="generar_excel",
        detalles={
            "archivo": str(ruta_destino),
            "filas": len(df),
            "columnas": len(columnas)
        },
        usuario=usuario,
        estado="EXITO"
    )

    return ruta_destino
